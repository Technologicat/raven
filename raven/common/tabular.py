"""Tables as files a person opens: TSV, Excel (`.xlsx`) and OpenDocument (`.ods`), chosen by the extension.

A report written for a person to read — an audit of what a tool removed, a worksheet for a reviewer to fill
in — is most useful in whatever that person's spreadsheet program opens best, so every such writer here
takes its format from the file name it is given, and every reader accepts all three.

The shape is the same in each: a header row of column names, then one row per record, plus optional comment
lines saying what the table is and how it was made. In a TSV the comments come first, each on a `# ` line;
in a spreadsheet they go on a second sheet, *Notes*, so that the first row of the first sheet is the header,
which is what sorting and filtering expect.
"""

from __future__ import annotations

__all__ = ["FORMATS", "format_of",
           "tsv_cell", "write_table", "read_table"]

import csv
import pathlib
import re
from collections.abc import Iterable

#: The formats `write_table` and `read_table` understand, by file extension.
FORMATS = ("tsv", "xlsx", "ods")

_DATA_SHEET = "Table"
_NOTES_SHEET = "Notes"

# Column widths in a spreadsheet, in characters: as wide as the widest cell, within these bounds. Wide
# enough that a title reads at a glance, narrow enough that one long abstract does not push the rest of the
# table off the screen.
_MIN_COLUMN_WIDTH = 8
_MAX_COLUMN_WIDTH = 60


def format_of(path: str | pathlib.Path) -> str:
    """The format of `path`, from its extension: one of `FORMATS`. Raises `ValueError` for anything else."""
    suffix = pathlib.Path(path).suffix.lower().lstrip(".")
    if suffix not in FORMATS:
        raise ValueError(f"format_of: '{path}': unknown table format '.{suffix}'; expected one of "
                         f"{', '.join('.' + fmt for fmt in FORMATS)}")
    return suffix


def tsv_cell(value: object) -> str:
    """`value` as one TSV cell: every run of whitespace, tabs and newlines included, collapsed to a space.

    A tab would end the cell and a newline the row, so a cell carrying either would misalign the file.
    """
    return re.sub(r"\s+", " ", str(value)).strip()


def write_table(path: str | pathlib.Path,
                comments: Iterable[str],
                columns: Iterable[str],
                rows: Iterable[Iterable[object]]) -> None:
    """Write a table to `path`, in the format its extension names (see `FORMATS`).

    `comments`: lines saying what the table is. A `# ` line each above the header in a TSV; one row each on
                a *Notes* sheet in a spreadsheet.
    `columns`: the column names, written as the header row.
    `rows`: one iterable of cells per row. A cell that is `None` is written empty. In a spreadsheet, an
            `int` or a `float` is written as a number, so that it sorts as one; anything else as text.

    In a TSV, every cell goes through `tsv_cell`, so each row is one line, and a cell containing a `"` is
    quoted with the `"` doubled — the CSV convention a spreadsheet's import follows, which otherwise reads
    a bare `"` as the start of a quoted cell and runs it on into the rows below.
    """
    fmt = format_of(path)
    comments = list(comments)
    columns = list(columns)
    rows = [["" if cell is None else cell for cell in row] for row in rows]
    if fmt == "tsv":
        _write_tsv(path, comments, columns, rows)
    elif fmt == "xlsx":
        _write_xlsx(path, comments, columns, rows)
    else:
        _write_ods(path, comments, columns, rows)


def read_table(path: str | pathlib.Path) -> list[dict[str, str]]:
    """Read a table written by `write_table`, or by hand in a spreadsheet program, from `path`.

    The format is taken from the extension (see `FORMATS`). The first row of the first sheet is the header;
    a TSV's `#` lines and a spreadsheet's other sheets are skipped. Rows that are wholly empty are skipped.

    Returns one dict per row, column name → cell, every cell as a string: empty for an empty cell, and a
    whole number without a decimal point, as a spreadsheet program shows it.
    """
    fmt = format_of(path)
    if fmt == "tsv":
        grid = _read_tsv(path)
    elif fmt == "xlsx":
        grid = _read_xlsx(path)
    else:
        grid = _read_ods(path)
    grid = [row for row in grid if any(cell.strip() for cell in row)]
    if not grid:
        return []
    header, *body = grid
    width = len(header)
    return [dict(zip(header, (row + [""] * width)[:width])) for row in body]


# --------------------------------------------------------------------------------
# TSV

def _write_tsv(path, comments, columns, rows) -> None:
    with open(path, "w", encoding="utf-8", newline="") as f:
        for comment in comments:
            f.write(f"# {tsv_cell(comment)}\n")
        writer = csv.writer(f, delimiter="\t", quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
        writer.writerow([tsv_cell(column) for column in columns])
        writer.writerows([tsv_cell(cell) for cell in row] for row in rows)


def _read_tsv(path) -> list[list[str]]:
    with open(path, encoding="utf-8", newline="") as f:
        lines = [line for line in f if not line.startswith("#")]
    return [list(row) for row in csv.reader(lines, delimiter="\t")]


# --------------------------------------------------------------------------------
# Spreadsheets

def _is_number(cell: object) -> bool:
    return isinstance(cell, (int, float)) and not isinstance(cell, bool)


def _column_widths(columns, rows) -> list[int]:
    widths = []
    for j, column in enumerate(columns):
        longest = max([len(str(column))] + [len(str(row[j])) for row in rows if j < len(row)])
        widths.append(max(_MIN_COLUMN_WIDTH, min(_MAX_COLUMN_WIDTH, longest + 2)))
    return widths


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _write_xlsx(path, comments, columns, rows) -> None:
    import xlsxwriter

    workbook = xlsxwriter.Workbook(str(path), {"strings_to_numbers": False,
                                               "strings_to_formulas": False,  # a title starting with `=` is a title
                                               "strings_to_urls": False})
    try:
        bold = workbook.add_format({"bold": True})
        sheet = workbook.add_worksheet(_DATA_SHEET)
        sheet.write_row(0, 0, columns, bold)
        for i, row in enumerate(rows, start=1):
            for j, cell in enumerate(row):
                if _is_number(cell):
                    sheet.write_number(i, j, cell)
                else:
                    sheet.write_string(i, j, str(cell))
        for j, width in enumerate(_column_widths(columns, rows)):
            sheet.set_column(j, j, width)
        sheet.freeze_panes(1, 0)
        if columns:
            sheet.autofilter(0, 0, len(rows), len(columns) - 1)
        if comments:
            notes = workbook.add_worksheet(_NOTES_SHEET)
            for i, comment in enumerate(comments):
                notes.write_string(i, 0, comment)
            notes.set_column(0, 0, _MAX_COLUMN_WIDTH * 2)
    finally:
        workbook.close()


def _read_xlsx(path) -> list[list[str]]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [[_cell_to_str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _write_ods(path, comments, columns, rows) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.style import Style, TableColumnProperties, TextProperties
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import P

    document = OpenDocumentSpreadsheet()
    bold = Style(name="header", family="table-cell")
    bold.addElement(TextProperties(fontweight="bold"))
    document.automaticstyles.addElement(bold)

    def text_cell(value: str, style=None) -> TableCell:
        cell = TableCell(valuetype="string", stylename=style) if style else TableCell(valuetype="string")
        cell.addElement(P(text=value))
        return cell

    def column_style(j: int, width: int) -> Style:
        style = Style(name=f"col{j}", family="table-column")
        # A character is roughly a quarter of a centimetre at the default font size.
        style.addElement(TableColumnProperties(columnwidth=f"{width * 0.22:.2f}cm"))
        document.automaticstyles.addElement(style)
        return style

    table = Table(name=_DATA_SHEET)
    for j, width in enumerate(_column_widths(columns, rows)):
        table.addElement(TableColumn(stylename=column_style(j, width)))
    header = TableRow()
    for column in columns:
        header.addElement(text_cell(str(column), bold))
    table.addElement(header)
    for row in rows:
        table_row = TableRow()
        for cell in row:
            if _is_number(cell):
                table_cell = TableCell(valuetype="float", value=cell)
                table_cell.addElement(P(text=str(cell)))
            else:
                table_cell = text_cell(str(cell))
            table_row.addElement(table_cell)
        table.addElement(table_row)
    document.spreadsheet.addElement(table)

    if comments:
        notes = Table(name=_NOTES_SHEET)
        for comment in comments:
            note_row = TableRow()
            note_row.addElement(text_cell(comment))
            notes.addElement(note_row)
        document.spreadsheet.addElement(notes)

    document.save(str(path))


def _read_ods(path) -> list[list[str]]:
    from odf.namespaces import TABLENS
    from odf.opendocument import load
    from odf.table import Table, TableCell, TableRow

    document = load(str(path))
    tables = document.spreadsheet.getElementsByType(Table)
    if not tables:
        return []

    def cell_text(cell) -> str:
        # A cell's text is its paragraphs, one per line.
        return "\n".join("".join(str(node) for node in p.childNodes)
                         for p in cell.childNodes if p.qname[1] == "p")

    grid = []
    for row in tables[0].getElementsByType(TableRow):
        # A spreadsheet program saves a run of identical cells, and of identical rows, as one element with a
        # repeat count — including the run of empty cells that pads every row out to the sheet's last column.
        # Expand the runs, then trim the empty tail they leave.
        cells = []
        for cell in row.getElementsByType(TableCell):
            repeat = int(cell.attributes.get((TABLENS, "number-columns-repeated"), 1))
            cells.extend([cell_text(cell)] * min(repeat, 1024))
        while cells and not cells[-1]:
            cells.pop()
        repeat = int(row.attributes.get((TABLENS, "number-rows-repeated"), 1))
        grid.extend([cells] * min(repeat, 1024) if cells else [[]])
    return grid
