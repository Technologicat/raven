"""Tests for raven.common.tabular: one table, written and read back as TSV, `.xlsx` and `.ods`."""

import csv
import pathlib

import pytest

from raven.common import tabular

COLUMNS = ("key", "title", "year", "score")
ROWS = [("a_2024", "Agents in the classroom", 2024, 0.75),
        ("b_2023", '"Hey ChatGPT": students ask', 2023, 1.5),
        ("c_2022", "=SUM(A1:A2) is a title, not a formula", 2022, None)]
COMMENTS = ["raven-test 1.0", "input: corpus.bib"]


def _expected_strings():
    """`ROWS` as `read_table` returns them: every cell a string, `None` empty, whole numbers without a `.0`."""
    def as_read(cell):
        if cell is None:
            return ""
        if isinstance(cell, float) and cell.is_integer():
            return str(int(cell))
        return str(cell)
    return [dict(zip(COLUMNS, (as_read(cell) for cell in row))) for row in ROWS]


class TestFormatOf:
    @pytest.mark.parametrize("name, fmt", [("a.tsv", "tsv"), ("a.XLSX", "xlsx"), ("dir.v2/a.ods", "ods")])
    def test_the_extension_names_the_format(self, name, fmt):
        assert tabular.format_of(name) == fmt

    def test_an_unknown_extension_is_refused_with_the_known_ones_named(self):
        with pytest.raises(ValueError, match=r"\.tsv, \.xlsx, \.ods"):
            tabular.format_of("report.csv")


class TestRoundTrip:
    """Whatever the format, the table that comes back is the table that went in."""

    @pytest.mark.parametrize("fmt", tabular.FORMATS)
    def test_rows_come_back_as_written(self, tmp_path, fmt):
        path = tmp_path / f"report.{fmt}"
        tabular.write_table(path, COMMENTS, COLUMNS, ROWS)
        assert tabular.read_table(path) == _expected_strings()

    @pytest.mark.parametrize("fmt", tabular.FORMATS)
    def test_an_empty_table_is_just_its_header(self, tmp_path, fmt):
        path = tmp_path / f"report.{fmt}"
        tabular.write_table(path, [], COLUMNS, [])
        assert tabular.read_table(path) == []


class TestSpreadsheets:
    def test_xlsx_numbers_are_numbers_and_comments_are_on_the_notes_sheet(self, tmp_path):
        import openpyxl

        path = tmp_path / "report.xlsx"
        tabular.write_table(path, COMMENTS, COLUMNS, ROWS)
        workbook = openpyxl.load_workbook(path)
        assert workbook.sheetnames == ["Table", "Notes"]
        table = workbook["Table"]
        assert [cell.value for cell in table[1]] == list(COLUMNS), "the header is the first row"
        assert table["C2"].data_type == "n" and table["C2"].value == 2024
        assert table["B4"].data_type == "s", "a title starting with `=` was written as a formula"
        assert [row[0].value for row in workbook["Notes"].iter_rows()] == COMMENTS

    def test_ods_numbers_are_numbers_and_comments_are_on_the_notes_sheet(self, tmp_path):
        from odf.opendocument import load
        from odf.table import Table, TableCell

        path = tmp_path / "report.ods"
        tabular.write_table(path, COMMENTS, COLUMNS, ROWS)
        tables = load(str(path)).spreadsheet.getElementsByType(Table)
        assert [table.getAttribute("name") for table in tables] == ["Table", "Notes"]
        year = tables[0].getElementsByType(TableCell)[len(COLUMNS) + 2]
        assert year.getAttribute("valuetype") == "float" and float(year.getAttribute("value")) == 2024

    def test_without_comments_there_is_no_notes_sheet(self, tmp_path):
        import openpyxl

        path = tmp_path / "report.xlsx"
        tabular.write_table(path, [], COLUMNS, ROWS)
        assert openpyxl.load_workbook(path).sheetnames == ["Table"]

    def test_an_ods_saved_by_a_spreadsheet_program_reads_back_whole(self, tmp_path):
        """LibreOffice saves a run of identical cells, or rows, as one element with a repeat count, and pads
        each row with empty cells out to the sheet's last column. Built here by hand in that shape."""
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        def cell(text=None, **repeat):
            element = TableCell(**repeat)
            if text is not None:
                element.addElement(P(text=text))
            return element

        document = OpenDocumentSpreadsheet()
        table = Table(name="Sheet1")
        for cells, rows_repeated in (([cell("key"), cell("verdict"), cell(numbercolumnsrepeated=1020)], 1),
                                     ([cell("a_2024"), cell("keep"), cell(numbercolumnsrepeated=1020)], 1),
                                     ([cell("x", numbercolumnsrepeated=2)], 2),
                                     ([cell(numbercolumnsrepeated=1022)], 1048000)):
            row = TableRow(numberrowsrepeated=rows_repeated) if rows_repeated > 1 else TableRow()
            for element in cells:
                row.addElement(element)
            table.addElement(row)
        document.spreadsheet.addElement(table)
        path = tmp_path / "filled-in.ods"
        document.save(str(path))

        assert tabular.read_table(path) == [{"key": "a_2024", "verdict": "keep"},
                                            {"key": "x", "verdict": "x"},
                                            {"key": "x", "verdict": "x"}]


class TestTsv:
    """A TSV report opens in a spreadsheet with one row per record and every cell intact."""

    def _read(self, path):
        """The data rows, parsed the way a spreadsheet's import does: `"` is a quote character."""
        lines = [line for line in path.read_text(encoding="utf-8").splitlines(keepends=True)
                 if not line.startswith("#")]
        return list(csv.reader(lines, delimiter="\t"))

    def test_a_cell_opening_with_a_quote_survives_the_round_trip(self, tmp_path):
        path = tmp_path / "report.tsv"
        rows = [("a_2024", '"Hey ChatGPT": students ask'), ("b_2024", "plain title")]
        tabular.write_table(path, ["tool 1.0"], ("key", "title"), rows)
        parsed = self._read(path)
        assert parsed[0] == ["key", "title"]
        assert [tuple(row) for row in parsed[1:]] == rows

    def test_a_quote_mid_cell_is_quoted_too(self, tmp_path):
        # Asserted on the text rather than by parsing it back: Python's reader tolerates a bare `"` inside a
        # cell, so a round trip passes whether or not it was quoted, while a spreadsheet import does not.
        path = tmp_path / "report.tsv"
        tabular.write_table(path, [], ("key", "title"), [("a_2024", 'The "learning assistant" in STEM')])
        assert path.read_text(encoding="utf-8").splitlines()[1] == 'a_2024\t"The ""learning assistant"" in STEM"'

    def test_tabs_and_newlines_in_a_cell_do_not_split_it(self, tmp_path):
        path = tmp_path / "report.tsv"
        tabular.write_table(path, [], ("key", "title"), [("a_2024", "two\tpart\ntitle ")])
        lines = [line for line in path.read_text(encoding="utf-8").splitlines() if not line.startswith("#")]
        assert lines == ["key\ttitle", "a_2024\ttwo part title"]

    def test_comments_come_first_as_hash_lines(self, tmp_path):
        path = tmp_path / "report.tsv"
        tabular.write_table(path, ["tool 1.0", "input: corpus.bib"], ("key",), [("a_2024",)])
        assert path.read_text(encoding="utf-8").splitlines() == ["# tool 1.0", "# input: corpus.bib",
                                                                 "key", "a_2024"]


class TestFilesSavedByLibreOffice:
    """Files a real spreadsheet program wrote, rather than files shaped by hand to look like them.

    Made by writing `ROWS` and `COMMENTS` with `write_table`, then converting with LibreOffice 24.2
    (`soffice --headless --convert-to`): the `.ods` from our `.xlsx`, the `.xlsx` from our `.ods`. So each
    was saved by a program other than the one that wrote the original, which is the situation a reviewer's
    filled-in sheet is in. Regenerate them the same way if `ROWS` changes.
    """

    @pytest.mark.parametrize("name", ["tabular_saved_by_libreoffice.ods", "tabular_saved_by_libreoffice.xlsx"])
    def test_reads_back_whole(self, name):
        path = pathlib.Path(__file__).parent / "data" / name
        assert tabular.read_table(path) == _expected_strings()

    def test_the_notes_sheet_survived_the_save(self):
        import openpyxl

        path = pathlib.Path(__file__).parent / "data" / "tabular_saved_by_libreoffice.xlsx"
        workbook = openpyxl.load_workbook(path)
        assert workbook.sheetnames == ["Table", "Notes"]
        assert [row[0].value for row in workbook["Notes"].iter_rows()] == COMMENTS
